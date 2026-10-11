"""Weekly openpyxl pack: KPIs, flagged SKUs/cities, reasons, trend."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.chart import BarChart, Reference

from app.config import BRAND, REPORTS_DIR
from app import engine

GOLD = "E89B1E"
TEAL = "0D9B7A"
ROSE = "E55A5A"
INK = "1A1D2E"
MUTED = "5A647A"
HEAD_FILL = PatternFill("solid", fgColor="1A1D2E")
GOLD_FILL = PatternFill("solid", fgColor=GOLD)
ROSE_FILL = PatternFill("solid", fgColor="FDECEC")
TEAL_FILL = PatternFill("solid", fgColor="E6F6F1")
THIN = Border(
    left=Side(style="thin", color="E2E6F0"),
    right=Side(style="thin", color="E2E6F0"),
    top=Side(style="thin", color="E2E6F0"),
    bottom=Side(style="thin", color="E2E6F0"),
)
HEAD_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE = Font(name="Calibri", bold=True, color=INK, size=16)
LABEL = Font(name="Calibri", color=MUTED, size=10)


def _header_row(ws, row: int, cols: int) -> None:
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN


def _write_df(ws, df, start_row: int = 1) -> int:
    rows = list(dataframe_to_rows(df, index=False, header=True))
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row, start=1):
            cell = ws.cell(row=start_row + r_i, column=c_i, value=val if val == val else None)
            cell.border = THIN
            cell.alignment = Alignment(vertical="center")
            if r_i == 0:
                cell.fill = HEAD_FILL
                cell.font = HEAD_FONT
                cell.alignment = Alignment(horizontal="center")
            elif "flag" in df.columns and c_i == list(df.columns).index("flag") + 1:
                if val is True:
                    cell.fill = ROSE_FILL
                    cell.font = Font(bold=True, color=ROSE)
    for i, col in enumerate(df.columns, start=1):
        ws.column_dimensions[chr(64 + i) if i < 27 else "A"].width = max(12, min(28, len(str(col)) + 6))
    return start_row + len(rows)


def build_weekly_pack(path: Path | None = None) -> Path:
    df = engine.joined()
    kpis = engine.kpis(df)
    skus = engine.flag_skus(df)
    cities = engine.flag_cities(df)
    pins = engine.flag_pincodes(df)
    reasons = engine.reasons_breakup(df)
    trend = engine.weekly_trend(df)

    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    cover["A1"] = f"{BRAND} — Returns & RTO weekly pack"
    cover["A1"].font = TITLE
    cover["A2"] = f"Generated {datetime.now().strftime('%d %b %Y %H:%M')} · SwiftLogic AI"
    cover["A2"].font = LABEL
    cover["A4"] = "KPI"
    cover["B4"] = "Value"
    _header_row(cover, 4, 2)
    items = [
        ("Orders", kpis["orders"]),
        ("GMV (₹)", kpis["gmv"]),
        ("RTO orders", kpis["rto_orders"]),
        ("RTO %", kpis["rto_pct"]),
        ("Returns", kpis["return_orders"]),
        ("Return %", kpis["return_pct"]),
        ("Failed GMV (₹)", kpis["failed_gmv"]),
        ("COD share %", kpis["cod_share"]),
        ("Flagged SKUs", int(skus.flag.sum())),
        ("Flagged cities", int(cities.flag.sum())),
    ]
    for i, (k, v) in enumerate(items, start=5):
        cover.cell(row=i, column=1, value=k).border = THIN
        cover.cell(row=i, column=2, value=v).border = THIN
    cover.column_dimensions["A"].width = 22
    cover.column_dimensions["B"].width = 18
    cover["A16"] = "Flags: SKU RTO ≥ 18% (≥40 orders) · City RTO ≥ 16% (≥50 orders) · Pin RTO ≥ 22% (≥20 orders)"
    cover["A16"].font = LABEL

    ws_sku = wb.create_sheet("High-RTO SKUs")
    _write_df(ws_sku, skus)

    ws_city = wb.create_sheet("High-RTO cities")
    _write_df(ws_city, cities)

    ws_pin = wb.create_sheet("Pincodes")
    _write_df(ws_pin, pins.head(40))

    ws_r = wb.create_sheet("Reasons")
    last = _write_df(ws_r, reasons)
    chart = BarChart()
    chart.title = "Failed orders by reason"
    chart.y_axis.title = "Orders"
    data = Reference(ws_r, min_col=4, min_row=1, max_row=last - 1)
    cats = Reference(ws_r, min_col=2, min_row=2, max_row=last - 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.style = 10
    ws_r.add_chart(chart, "G2")

    ws_t = wb.create_sheet("Weekly trend")
    _write_df(ws_t, trend)

    stamp = datetime.now().strftime("%Y%m%d")
    out = path or (REPORTS_DIR / f"rto_weekly_pack_{stamp}.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
